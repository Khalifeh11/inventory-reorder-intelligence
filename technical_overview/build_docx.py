"""Generate Atlas_Technical_Overview.docx (1-2 pages) from the markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

SRC = Path(__file__).resolve().parent / "Atlas_Technical_Overview.md"
OUT = Path(__file__).resolve().parent / "Atlas_Technical_Overview.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GRAY = RGBColor(0x5A, 0x6A, 0x7A)


def _set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _add_inline(para, text: str, *, bold=False, italic=False, mono=False, color=None, size_pt=10):
    # Resolve **bold**, *italic*, `code` inside text if it looks raw
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        run = para.add_run()
        run_bold = bold
        run_italic = italic
        run_mono = mono
        t = tok
        if tok.startswith("**") and tok.endswith("**"):
            t = tok[2:-2]
            run_bold = True
        elif tok.startswith("*") and tok.endswith("*") and not tok.startswith("**"):
            t = tok[1:-1]
            run_italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            t = tok[1:-1]
            run_mono = True
        run.text = t
        run.bold = run_bold
        run.italic = run_italic
        run.font.size = Pt(size_pt)
        run.font.name = "Menlo" if run_mono else "Calibri"
        if color is not None:
            run.font.color.rgb = color


def _set_margins(doc):
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)


def build() -> Path:
    md = SRC.read_text()
    doc = Document()
    _set_margins(doc)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    lines = md.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Menlo"
                r.font.size = Pt(8.5)
                r.font.color.rgb = NAVY
                in_code = False
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            _add_inline(p, line[2:], bold=True, color=NAVY, size_pt=18)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            _add_inline(p, line[3:], bold=True, color=NAVY, size_pt=13)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            _add_inline(p, line[4:], bold=True, color=ORANGE, size_pt=11)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        if line.strip() == "---":
            p = doc.add_paragraph()
            r = p.add_run("_" * 80)
            r.font.color.rgb = GRAY
            r.font.size = Pt(8)
            i += 1
            continue

        # Tables
        if line.strip().startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            rows = []
            for ln in tbl_lines:
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                if all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_i, row in enumerate(rows):
                    for c_i, cell_text in enumerate(row):
                        cell = table.rows[r_i].cells[c_i]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        if r_i == 0:
                            _add_inline(p, cell_text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size_pt=9)
                            _set_cell_shading(cell, "1F2A44")
                        else:
                            _add_inline(p, cell_text, color=NAVY, size_pt=9)
                            if r_i % 2 == 0:
                                _set_cell_shading(cell, "F4F6F8")
            continue

        if re.match(r"^\s*[-*] ", line):
            while i < len(lines) and re.match(r"^\s*[-*] ", lines[i]):
                txt = re.sub(r"^\s*[-*] ", "", lines[i])
                p = doc.add_paragraph(style="List Bullet")
                _add_inline(p, txt, color=NAVY, size_pt=10)
                i += 1
            continue

        if re.match(r"^\s*\d+\.\s", line):
            while i < len(lines) and re.match(r"^\s*\d+\.\s", lines[i]):
                txt = re.sub(r"^\s*\d+\.\s", "", lines[i])
                p = doc.add_paragraph(style="List Number")
                _add_inline(p, txt, color=NAVY, size_pt=10)
                i += 1
            continue

        if line.strip() == "":
            i += 1
            continue

        # Plain paragraph (merge consecutive non-block lines)
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not _is_block_start(lines[i]):
            para_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        _add_inline(p, " ".join(para_lines), color=NAVY, size_pt=10)

    doc.save(OUT)
    return OUT


def _is_block_start(line: str) -> bool:
    if not line:
        return True
    if line.startswith("#") or line.startswith("|") or line.startswith("```"):
        return True
    if line.strip() == "---":
        return True
    if re.match(r"^\s*[-*] ", line):
        return True
    if re.match(r"^\s*\d+\.\s", line):
        return True
    return False


if __name__ == "__main__":
    path = build()
    print(f"Built {path} ({path.stat().st_size // 1024} KB)")
